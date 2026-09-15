"""知识库文档服务（RAG 服务内部）：上传存储、文本解析、分块、向量化入库、检索与管理。

文件存储策略：
- 文件保存到 settings.upload_dir 下（本服务唯一持有磁盘写入）。
- 磁盘上文件名使用 uuid，避免重名/路径穿越；原始文件名仅记录在 DB。
- DB 中 file_path 为相对 upload_dir 的路径，便于整体迁移。
"""

import logging
import os
import uuid
from datetime import datetime
from uuid import NAMESPACE_OID, uuid5

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.config import settings
from app.exceptions import AppError, NotFoundError
from app.models import Document, DocumentChunk
from app.services import chunking_service, embedding_service, qdrant_service

logger = logging.getLogger(__name__)

# 支持解析的扩展名 → 解析器
ALLOWED_EXTS = {".pdf", ".docx", ".txt", ".md"}

# 文档类型标签
DOC_TYPES = {
    "guide": "临床指南",
    "literature": "医学文献",
    "drug": "药品说明书",
    "case": "疑难病例",
    "norm": "院内规范",
    "other": "其他",
}


def _hospital_upload_dir(hospital_id: int) -> str:
    """医院专属上传目录（绝对路径）。"""
    base = os.path.abspath(settings.upload_dir)
    d = os.path.join(base, f"hospital_{hospital_id}")
    os.makedirs(d, exist_ok=True)
    return d


def _save_upload_file(hospital_id: int, file_name: str, content: bytes) -> str:
    """保存文件到上传目录，返回相对 upload_dir 的路径（POSIX 风格）。"""
    ext = os.path.splitext(file_name)[1].lower()
    stored_name = f"{uuid.uuid4().hex}{ext}"
    rel_path = f"hospital_{hospital_id}/{stored_name}"
    abs_path = os.path.join(os.path.abspath(settings.upload_dir), rel_path)
    os.makedirs(os.path.dirname(abs_path), exist_ok=True)
    with open(abs_path, "wb") as f:
        f.write(content)
    return rel_path


def _abs_file_path(rel_path: str) -> str:
    return os.path.join(os.path.abspath(settings.upload_dir), rel_path)


# ---------- 文本解析 ----------

def extract_text(file_path: str, ext: str) -> str:
    """按扩展名抽取纯文本。

    PDF 损坏/加密（或缺少 cryptography 依赖）时抛 AppError，
    让调用方能给出明确原因，而不是笼统的“索引异常”。
    """
    if ext == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(file_path)
        except Exception as exc:  # noqa: BLE001 - 加密 / 损坏 / 缺依赖
            raise AppError(
                f"PDF 解析失败（文档可能已加密或已损坏）：{exc}",
                code="MED_DOC_PARSE_FAILED",
            ) from exc
        try:
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:  # noqa: BLE001 - 单页解密/解析失败
            raise AppError(
                f"PDF 正文提取失败（文档可能已加密）：{exc}",
                code="MED_DOC_PARSE_FAILED",
            ) from exc
        return "\n".join(pages)
    if ext == ".docx":
        import docx

        try:
            d = docx.Document(file_path)
        except Exception as exc:  # noqa: BLE001 - 损坏或非法的 docx
            raise AppError(
                f"Word 文档解析失败（文件可能已损坏）：{exc}",
                code="MED_DOC_PARSE_FAILED",
            ) from exc
        return "\n".join(p.text for p in d.paragraphs if p.text and p.text.strip())
    if ext in (".txt", ".md"):
        for enc in ("utf-8", "gb18030"):
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    return ""


def chunk_text(text: str) -> list[chunking_service.Chunk]:
    """分块入口：优先 LLM 语义切分，不可用时回退规则切分。"""
    if settings.llm_chunking:
        try:
            chunks = chunking_service.llm_chunk_text(text)
            if chunks:
                return chunks
        except Exception:  # noqa: BLE001 - 切分异常一律回退
            logger.exception("LLM 切分异常，回退规则切分")
    return chunking_service.split_text(text)


# ---------- 入库流程 ----------

def _build_chunk_records(
    document_id: int, hospital_id: int, chunks: list[chunking_service.Chunk]
) -> list[tuple[DocumentChunk, str, list[float]]]:
    """向量化分块并生成 chunk 记录（DB 行 + Qdrant 点）。

    Qdrant 点 ID 使用确定性 UUID（Qdrant local 模式要求合法 UUID）。
    """
    vectors = embedding_service.embed_texts([c.content for c in chunks])
    records: list[tuple[DocumentChunk, str, list[float]]] = []
    for idx, (ch, vector) in enumerate(zip(chunks, vectors)):
        point_id = str(uuid5(NAMESPACE_OID, f"doc{document_id}-c{idx}"))
        chunk = DocumentChunk(
            hospital_id=hospital_id,
            document_id=document_id,
            chunk_index=idx,
            title=ch.title or None,
            content=ch.content,
            char_count=len(ch.content),
            point_id=point_id,
        )
        records.append((chunk, point_id, vector))
    return records


def create_document(
    db: Session,
    *,
    hospital_id: int,
    file_name: str,
    file_content: bytes,
    title: str,
    doc_type: str,
    sub_type: str | None,
    source: str,
    source_type: str = "system",
    remark: str | None,
    created_by: str,
) -> Document:
    """保存文件 → 建文档记录（status=parsing）。立即返回，索引由后台任务执行。"""
    ext = os.path.splitext(file_name)[1].lower()
    if ext not in ALLOWED_EXTS:
        raise AppError(
            f"暂不支持 {ext} 格式，仅支持 {' / '.join(sorted(ALLOWED_EXTS))}",
            code="MED_DOC_TYPE_UNSUPPORTED",
        )
    if len(file_content) > settings.max_upload_mb * 1024 * 1024:
        raise AppError(
            f"文件超过 {settings.max_upload_mb}MB 上限", code="MED_DOC_TOO_LARGE"
        )

    rel_path = _save_upload_file(hospital_id, file_name, file_content)

    doc = Document(
        hospital_id=hospital_id,
        title=title or os.path.splitext(file_name)[0],
        file_name=file_name,
        file_ext=ext,
        file_size=len(file_content),
        file_path=rel_path,
        doc_type=doc_type if doc_type in DOC_TYPES else "other",
        sub_type=(sub_type or "").strip() or None,
        source=source or "内部上传",
        source_type=source_type if source_type in ("system", "ima") else "system",
        remark=remark,
        status="parsing",
        created_by=created_by,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def create_ima_placeholder(
    db: Session,
    *,
    hospital_id: int,
    media_id: str,
    title: str,
    file_ext: str = ".md",
    doc_type: str = "other",
    source: str = "腾讯ima",
    remark: str | None = None,
    created_by: str = "system",
) -> Document:
    """为 IMA 文件建立占位记录：只登记元信息，不下载正文、不切分、不向量化。

    两步式导入的第一步，秒返回。第二步由 index_document_text 用 media_id
    取回内容后执行索引。
    """
    media_id = (media_id or "").strip()
    if not media_id:
        raise AppError("media_id 不能为空", code="MED_DOC_MEDIA_ID_REQUIRED")

    # 同一医院内同一 media 不重复登记
    existing = db.scalar(
        select(Document).where(
            Document.hospital_id == hospital_id,
            Document.source_media_id == media_id,
            Document.deleted_at.is_(None),
        )
    )
    if existing is not None:
        return existing

    file_name = f"{title or media_id}{file_ext}"
    doc = Document(
        hospital_id=hospital_id,
        title=title or file_name,
        file_name=file_name,
        file_ext=file_ext,
        file_size=0,
        # 未落盘前用 ima:// 标识，索引时写入文件后替换为真实相对路径
        file_path=f"ima://{media_id}",
        doc_type=doc_type if doc_type in DOC_TYPES else "other",
        source=source,
        source_type="ima",
        source_media_id=media_id,
        remark=remark,
        status="uploaded",
        created_by=created_by,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def index_document_text(
    db: Session,
    doc: Document,
    text: str,
    *,
    file_ext: str | None = None,
) -> Document:
    """用外部提供的正文为文档建立索引（两步式导入的第二步）。

    正文由服务 A 从 IMA 取回后传入；本服务负责落盘、切分、向量化。

    注意：服务 A 在触发索引时已把状态置为 parsing，本方法不再拒绝 parsing
    （防重由服务 A 的 start_document_index 负责），否则两步式流程会被自己挡住。
    """
    # 重新索引需先清理旧数据
    qdrant_service.delete_document_points(doc.hospital_id, doc.id)
    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).delete()

    ext = (file_ext or doc.file_ext or ".md").lower()
    content = (text or "").strip()
    if not content:
        doc.status = "failed"
        doc.error_message = "未能从 IMA 获取到可索引的内容"
        doc.chunk_count = 0
        doc.vector_count = 0
        db.commit()
        db.refresh(doc)
        return doc

    # 原文落盘，便于后续重建索引/复用，同时把占位路径替换为真实路径
    try:
        raw = content.encode("utf-8")
        rel_path = _save_upload_file(doc.hospital_id, f"{doc.title}{ext}", raw)
        doc.file_path = rel_path
        doc.file_ext = ext
        doc.file_size = len(raw)
        db.commit()
    except OSError as exc:
        logger.warning("IMA 文档落盘失败 doc_id=%s: %s", doc.id, exc)

    doc.status = "parsing"
    doc.error_message = None
    db.commit()

    try:
        chunks = chunk_text(content)
        if not chunks:
            raise AppError("文本内容为空", code="MED_DOC_EMPTY_TEXT")

        if settings.llm_chunking and not doc.summary:
            doc.summary = chunking_service.summarize_text(content)

        records = _build_chunk_records(doc.id, doc.hospital_id, chunks)
        qdrant_service.upsert_chunks(
            doc.hospital_id,
            doc.id,
            [
                {"point_id": point_id, "content": _chunk.content, "vector": vector}
                for _chunk, point_id, vector in records
            ],
        )
        db.add_all([r[0] for r in records])
        doc.chunk_count = len(records)
        doc.vector_count = len(records)
        doc.status = "ready"
        doc.error_message = None
    except (embedding_service.EmbeddingError, qdrant_service.VectorDbError) as exc:
        doc.status = "failed"
        doc.error_message = str(exc)[:500]
    except Exception as exc:  # noqa: BLE001 - 一律置 failed，避免逃逸成 HTTP 500
        logger.exception("IMA 文档索引失败 doc_id=%s", doc.id)
        doc.status = "failed"
        doc.error_message = f"索引失败：{exc}"[:500]

    db.commit()
    db.refresh(doc)
    return doc


def run_document_indexing(doc_id: int, hospital_id: int) -> None:
    """后台索引文档：独立会话加载记录并执行解析 → 分块 → 向量化 → 入库。

    供上传/搬运接口在响应返回后调用（FastAPI BackgroundTasks + 线程池），
    不依赖请求级会话，任何异常统一置为 failed，不向上抛出。
    """
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        doc = db.scalar(
            select(Document).where(
                Document.id == doc_id,
                Document.hospital_id == hospital_id,
                Document.deleted_at.is_(None),
            )
        )
        if doc is None:
            return
        try:
            _process_document(db, doc)
        except (embedding_service.EmbeddingError, qdrant_service.VectorDbError) as exc:
            doc.status = "failed"
            doc.error_message = str(exc)[:500]
            db.commit()
    except Exception as exc:  # noqa: BLE001 - 兜底：后台任务不向外抛
        logger.exception("后台索引文档 %s 失败", doc_id)
        try:
            doc = db.get(Document, doc_id)
            if doc is not None and doc.status != "failed":
                doc.status = "failed"
                # 记录真实原因，便于定位（此前固定文案掩盖了加密 PDF、缺依赖等问题）
                doc.error_message = f"索引失败：{exc}"[:500]
                db.commit()
        except Exception:  # noqa: BLE001
            db.rollback()
    finally:
        db.close()


def _process_document(db: Session, doc: Document) -> None:
    """解析文本 → 分块 → 向量化 → 写 Qdrant → 写 chunk 表。"""
    abs_path = _abs_file_path(doc.file_path)
    text = extract_text(abs_path, doc.file_ext)
    if not text or not text.strip():
        doc.status = "failed"
        doc.error_message = "未从文件中提取到文本内容"
        db.commit()
        return

    chunks = chunk_text(text)
    if not chunks:
        doc.status = "failed"
        doc.error_message = "文本内容为空"
        db.commit()
        return

    # LLM 生成文档摘要（失败不中断上传）
    if settings.llm_chunking and not doc.summary:
        doc.summary = chunking_service.summarize_text(text)

    records = _build_chunk_records(doc.id, doc.hospital_id, chunks)
    qdrant_service.upsert_chunks(
        doc.hospital_id,
        doc.id,
        [
            {"point_id": point_id, "content": _chunk.content, "vector": vector}
            for _chunk, point_id, vector in records
        ],
    )

    db.add_all([r[0] for r in records])
    doc.chunk_count = len(records)
    doc.vector_count = len(records)
    doc.status = "ready"
    doc.error_message = None
    db.commit()
    db.refresh(doc)


def reindex_document(db: Session, doc: Document) -> Document:
    """清空旧索引后重新解析入库（模型/解析升级后可重跑）。"""
    qdrant_service.delete_document_points(doc.hospital_id, doc.id)
    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).delete()
    doc.chunk_count = 0
    doc.vector_count = 0
    doc.status = "parsing"
    doc.error_message = None
    db.commit()

    try:
        _process_document(db, doc)
    except (embedding_service.EmbeddingError, qdrant_service.VectorDbError) as exc:
        doc.status = "failed"
        doc.error_message = str(exc)[:500]
        db.commit()
        db.refresh(doc)
    except Exception as exc:  # noqa: BLE001 - 同步路径：一律置 failed，避免逃逸成 HTTP 500
        logger.exception("重建索引失败 doc_id=%s", doc.id)
        doc.status = "failed"
        doc.error_message = f"重建索引失败：{exc}"[:500]
        db.commit()
        db.refresh(doc)
    return doc


# ---------- 查询 ----------

def list_documents(
    db: Session,
    *,
    hospital_id: int,
    keyword: str | None = None,
    doc_type: str | None = None,
    status: str | None = None,
    source_type: str | None = None,
    page: int = 1,
    page_size: int = 10,
) -> tuple[list[Document], int]:
    stmt = select(Document).where(
        Document.hospital_id == hospital_id, Document.deleted_at.is_(None)
    )
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(Document.title.like(like), Document.file_name.like(like)))
    if doc_type:
        stmt = stmt.where(Document.doc_type == doc_type)
    if status:
        stmt = stmt.where(Document.status == status)
    if source_type:
        stmt = stmt.where(Document.source_type == source_type)

    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    stmt = (
        stmt.order_by(Document.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    return list(db.scalars(stmt).all()), total


def get_overview(db: Session, *, hospital_id: int) -> dict:
    """知识库总览统计。"""
    base = select(Document).where(
        Document.hospital_id == hospital_id, Document.deleted_at.is_(None)
    )
    total = db.scalar(select(func.count()).select_from(base.subquery())) or 0

    def _count(status: str) -> int:
        return (
            db.scalar(
                select(func.count())
                .select_from(base.where(Document.status == status).subquery())
            )
            or 0
        )

    ready = _count("ready")
    parsing = _count("parsing")
    failed = _count("failed")

    total_chunks = (
        db.scalar(
            select(func.coalesce(func.sum(Document.chunk_count), 0)).where(
                Document.hospital_id == hospital_id, Document.deleted_at.is_(None)
            )
        )
        or 0
    )
    total_size = (
        db.scalar(
            select(func.coalesce(func.sum(Document.file_size), 0)).where(
                Document.hospital_id == hospital_id, Document.deleted_at.is_(None)
            )
        )
        or 0
    )

    type_rows = db.execute(
        select(Document.doc_type, func.count())
        .where(Document.hospital_id == hospital_id, Document.deleted_at.is_(None))
        .group_by(Document.doc_type)
    ).all()
    doc_type_distribution = [
        {"doc_type": t, "label": DOC_TYPES.get(t, t), "count": c}
        for t, c in type_rows
    ]

    recent = list(
        db.scalars(
            base.order_by(Document.id.desc()).limit(5)
        ).all()
    )

    return {
        "total_documents": total,
        "ready_documents": ready,
        "parsing_documents": parsing,
        "failed_documents": failed,
        "total_chunks": int(total_chunks),
        "total_vectors": qdrant_service.count_vectors(),
        "total_size_bytes": int(total_size),
        "doc_type_distribution": doc_type_distribution,
        "status_distribution": [
            {"status": "ready", "label": "已索引", "count": ready},
            {"status": "parsing", "label": "处理中", "count": parsing},
            {"status": "failed", "label": "失败", "count": failed},
        ],
        "recent_documents": recent,
        "embedding_model": settings.embedding_model_name,
        "embedding_ready": embedding_service.embedding_ready(),
        "vector_db_connected": qdrant_service.vector_db_ready(),
        "upload_dir": os.path.abspath(settings.upload_dir),
    }


def get_document_or_404(db: Session, hospital_id: int, doc_id: int) -> Document:
    doc = db.scalar(
        select(Document).where(
            Document.id == doc_id,
            Document.hospital_id == hospital_id,
            Document.deleted_at.is_(None),
        )
    )
    if doc is None:
        raise NotFoundError("文档不存在或已删除")
    return doc


def get_download_file_path(doc: Document) -> str | None:
    """返回可下载的原始文件绝对路径；IMA 占位（ima://）或文件已丢失时返回 None。"""
    rel = doc.file_path or ""
    if not rel or rel.startswith("ima://"):
        return None
    abs_path = _abs_file_path(rel)
    if not os.path.exists(abs_path) or not os.path.isfile(abs_path):
        return None
    return abs_path


def get_document_detail(db: Session, doc: Document) -> dict:
    """文档详情：元信息 + 分块列表（含标题/字数，便于前端展示）。"""
    chunks = list(
        db.scalars(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == doc.id)
            .order_by(DocumentChunk.chunk_index)
        ).all()
    )
    return {
        "document": doc,
        "chunks": [
            {
                "id": c.id,
                "chunk_index": c.chunk_index,
                "title": c.title,
                "content": c.content,
                "char_count": c.char_count,
            }
            for c in chunks
        ],
    }


def delete_document(db: Session, doc: Document) -> None:
    """删除文档：清向量 → 清分块 → 删文件 → 软删记录。"""
    try:
        qdrant_service.delete_document_points(doc.hospital_id, doc.id)
    except Exception as exc:  # pragma: no cover
        logger.warning("删除文档 %s 向量失败：%s", doc.id, exc)
    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).delete()
    try:
        abs_path = _abs_file_path(doc.file_path)
        if os.path.exists(abs_path):
            os.remove(abs_path)
    except OSError as exc:  # pragma: no cover
        logger.warning("删除文件失败 %s：%s", doc.file_path, exc)
    doc.deleted_at = datetime.now()
    db.commit()


# ---------- 向量检索 ----------

def search_knowledge(
    db: Session, *, hospital_id: int, query: str, limit: int = 10
) -> dict:
    """语义检索：query → 向量 → Qdrant topK → 回查文档元信息。"""
    from app.schemas.knowledge import SearchHit

    if not query.strip():
        raise AppError("检索关键词不能为空", code="MED_SEARCH_EMPTY")

    vector = embedding_service.embed_texts([query.strip()])[0]
    hits = qdrant_service.search(hospital_id, vector, limit=limit * 2)

    doc_ids = {h["document_id"] for h in hits}
    docs: dict[int, Document] = {}
    if doc_ids:
        rows = db.scalars(
            select(Document).where(Document.id.in_(doc_ids), Document.deleted_at.is_(None))
        ).all()
        docs = {d.id: d for d in rows}

    # 过滤"幽灵向量"：向量库中存在但 DB 已删除/不存在的文档，避免返回脏数据
    items = [
        SearchHit(
            document_id=h["document_id"],
            title=d.title,
            file_name=d.file_name,
            doc_type=d.doc_type,
            content=h["content"],
            score=h["score"],
        )
        for h in hits
        if (d := docs.get(h["document_id"])) is not None
    ]
    return {"query": query, "hits": items[:limit], "total": len(items)}
