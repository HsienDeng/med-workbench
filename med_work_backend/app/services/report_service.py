"""患者诊疗分析报告（Word / PDF）生成服务。

数据装配（build_bundle）与排版渲染（render_docx / render_pdf）分离：
- bundle 是纯字典，两种格式共用同一份数据，排版互不影响；
- docx 用 python-docx 生成，交付物为医生可编辑再打印的 .docx；
- pdf 用 reportlab 生成，交付物为版式固定、便于打印归档的 .pdf。
"""
from __future__ import annotations

import io
import logging
import os
from datetime import date, datetime
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape as _xml_escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AnalysisRecord, MedicalRecord, Patient, RbacUser
from app.services import dictionary_service, timeline_service

logger = logging.getLogger(__name__)

_FONT = "微软雅黑"
_PRIMARY = RGBColor(0x2D, 0x6C, 0xDF)
_MUTED = RGBColor(0x8A, 0x94, 0xA6)
_DISCLAIMER = (
    "本报告由 MedAI Workbench 基于院内数据与 AI 模型自动生成，"
    "仅供临床参考，最终诊疗决策以主管医师判断为准。"
)


def _dict_label(db: Session, hospital_id: int, dict_code: str, item_code: str) -> str:
    """字典项编码 → 显示名；字典缺失或未命中时回退编码本身。"""
    if not item_code:
        return "-"
    try:
        _, items = dictionary_service.get_options(db, hospital_id=hospital_id, dict_code=dict_code)
    except Exception as exc:  # noqa: BLE001
        logger.warning("报告字典翻译失败 dict=%s: %s", dict_code, exc)
        return item_code
    for item in items:
        if item.item_code == item_code:
            return item.item_label
    return item_code


def _fmt_time(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M") if value else "-"


def build_bundle(
    db: Session,
    patient: Patient,
    operator: RbacUser,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
) -> dict:
    """装配报告所需的全部数据（患者信息 + 病历 + AI 分析）。"""
    hospital_id = patient.hospital_id

    patient_info: list[tuple[str, str]] = [
        ("患者编号", patient.patient_no or "-"),
        ("姓名", patient.name),
        ("性别", timeline_service.GENDER_LABELS.get(patient.gender, patient.gender or "-")),
        ("年龄", f"{patient.age} 岁" if patient.age else "-"),
        ("身高", f"{patient.height} cm" if patient.height else "-"),
        ("体重", f"{patient.weight} kg" if patient.weight else "-"),
        ("BMI", str(patient.bmi) if patient.bmi else "-"),
        ("腰围", f"{patient.waistline} cm" if patient.waistline else "-"),
        ("手机号", patient.phone or "-"),
        ("科室", _dict_label(db, hospital_id, "department", patient.dept)),
        ("状态", _dict_label(db, hospital_id, "patient_status", patient.status)),
        ("建档时间", _fmt_time(patient.created_at)),
    ]

    records = db.scalars(
        select(MedicalRecord)
        .where(
            MedicalRecord.hospital_id == hospital_id,
            MedicalRecord.patient_id == patient.id,
            MedicalRecord.deleted_at.is_(None),
            *timeline_service.time_range_conditions(MedicalRecord.created_at, start_date, end_date),
        )
        .order_by(MedicalRecord.created_at.desc())
    ).all()
    record_items = [
        {
            "id": record.id,
            "created_at": _fmt_time(record.created_at),
            "fields": timeline_service.record_fields(record),
        }
        for record in records
    ]

    analysis_rows = db.execute(
        select(AnalysisRecord, RbacUser.real_name, RbacUser.username)
        .outerjoin(RbacUser, RbacUser.id == AnalysisRecord.user_id)
        .where(
            AnalysisRecord.hospital_id == hospital_id,
            AnalysisRecord.patient_id == patient.id,
            *timeline_service.time_range_conditions(AnalysisRecord.created_at, start_date, end_date),
        )
        .order_by(AnalysisRecord.created_at.desc())
    ).all()
    analysis_items: list[dict] = []
    for row, real_name, username in analysis_rows:
        result = row.result or {}
        analysis_items.append(
            {
                "id": row.id,
                "created_at": _fmt_time(row.created_at),
                "type_label": timeline_service.ANALYSIS_TYPE_LABELS.get(row.analysis_type, row.analysis_type),
                "model": row.model or "-",
                "status": row.status,
                "operator": real_name or username or "-",
                "error": row.error or "",
                "summary": result.get("summary") or {},
                "attention": result.get("attention") or [],
                "evidence": result.get("evidence") or [],
            }
        )

    if start_date or end_date:
        range_label = f"{start_date or '最早'} ~ {end_date or '至今'}"
    else:
        range_label = "全部记录"

    return {
        "patient": {
            "name": patient.name,
            "patient_no": patient.patient_no,
            "primary_diag": patient.primary_diag or "",
            "info": patient_info,
        },
        "records": record_items,
        "analyses": analysis_items,
        "range_label": range_label,
        "generated_at": datetime.now(),
        "operator": operator.real_name or operator.username,
    }


# ------------------------- docx 排版 -------------------------


def _apply_font(run, size: float | None = None, bold: bool = False, color: RGBColor | None = None) -> None:
    """设置字体（含中文 eastAsia，避免 Word 回落到非预期字体）。"""
    run.font.name = _FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), _FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def _setup(document: Document) -> None:
    """正文默认字体（中文字体需在 Normal 样式上设置 eastAsia）。"""
    normal = document.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(10.5)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), _FONT)


def _heading(document: Document, text: str, size: float = 14) -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after = Pt(6)
    _apply_font(para.add_run(text), size=size, bold=True, color=_PRIMARY)


def _body(document: Document, text: str, size: float = 10.5, muted: bool = False) -> None:
    para = document.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    _apply_font(para.add_run(text), size=size, color=_MUTED if muted else None)


def _key_value(document: Document, label: str, value: str) -> None:
    """「标签：值」段落，标签加粗、值换行友好（长文本保留原换行）。"""
    para = document.add_paragraph()
    para.paragraph_format.space_after = Pt(2)
    _apply_font(para.add_run(f"{label}："), size=10.5, bold=True)
    _apply_font(para.add_run(str(value) if value else "-"), size=10.5)


def _info_table(document: Document, info: list[tuple[str, str]]) -> None:
    """患者信息两列对齐表格：每行两组「标签 | 值」。"""
    table = document.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    for pair in [info[i : i + 2] for i in range(0, len(info), 2)]:
        cells = table.add_row().cells
        for index, (label, value) in enumerate(pair):
            _apply_font(cells[index * 2].paragraphs[0].add_run(label), size=10, bold=True)
            _apply_font(cells[index * 2 + 1].paragraphs[0].add_run(str(value)), size=10)
    if len(info) % 2 == 1:
        cells = table.rows[-1].cells
        for index in (2, 3):
            _apply_font(cells[index].paragraphs[0].add_run(""), size=10)


def _render_records(document: Document, records: list[dict]) -> None:
    _heading(document, f"二、病历记录（{len(records)} 份）")
    if not records:
        _body(document, "所选时间范围内暂无病历记录。", muted=True)
        return
    for index, item in enumerate(records, start=1):
        _heading(document, f"病历 {index} · 记录于 {item['created_at']}", size=11.5)
        if not item["fields"]:
            _body(document, "（该病历暂无填写内容）", muted=True)
            continue
        for label, value in item["fields"]:
            _key_value(document, label, value)


def _render_analyses(document: Document, analyses: list[dict]) -> None:
    _heading(document, f"三、AI 分析报告（{len(analyses)} 次）")
    if not analyses:
        _body(document, "所选时间范围内暂无 AI 分析记录。", muted=True)
        return
    for index, item in enumerate(analyses, start=1):
        _heading(document, f"分析 {index} · {item['type_label']} · {item['created_at']}", size=11.5)
        _body(
            document,
            f"模型：{item['model']}　发起人：{item['operator']}　"
            f"状态：{'成功' if item['status'] == 'done' else '失败'}",
            muted=True,
        )
        if item["status"] != "done":
            _body(document, f"失败原因：{item['error'] or '未记录'}")
            continue

        summary = item["summary"] or {}
        if summary:
            _body(document, "综合结论")
            for label, value in summary.items():
                _key_value(document, str(label), str(value))
        else:
            _body(document, "综合结论：（无）", muted=True)

        if item["attention"]:
            _body(document, "关注点")
            for point in item["attention"]:
                level = str(point.get("level") or "").strip()
                title = str(point.get("title") or "").strip()
                desc = str(point.get("description") or "").strip()
                label = f"【{level}】{title}".strip() if level else (title or "关注点")
                _key_value(document, label, desc)

        if item["evidence"]:
            _body(document, "循证依据")
            for ev_index, ev in enumerate(item["evidence"], start=1):
                source = str(ev.get("source") or "未命名来源")
                relevance = ev.get("relevance")
                if relevance is not None:
                    source = f"{source}（相关度 {relevance}）"
                _key_value(document, f"依据 {ev_index}", source)


def render_docx(bundle: dict) -> bytes:
    """把 bundle 渲染为 .docx 文件字节流。"""
    document = Document()
    _setup(document)

    patient = bundle["patient"]
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    _apply_font(title.add_run("患者诊疗分析报告"), size=20, bold=True)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _apply_font(
        sub.add_run(
            f"{patient['name']}　{patient['patient_no']}　"
            f"统计范围：{bundle['range_label']}　生成时间：{_fmt_time(bundle['generated_at'])}"
        ),
        size=9.5,
        color=_MUTED,
    )

    _heading(document, "一、患者基本信息")
    _info_table(document, patient["info"])
    if patient["primary_diag"]:
        _key_value(document, "主诊断", patient["primary_diag"])

    _render_records(document, bundle["records"])
    _render_analyses(document, bundle["analyses"])

    footer = document.add_paragraph()
    footer.paragraph_format.space_before = Pt(18)
    _apply_font(footer.add_run(_DISCLAIMER), size=9, color=_MUTED)
    _apply_font(
        footer.add_run(f"　导出人：{bundle['operator']}"),
        size=9,
        color=_MUTED,
    )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_filename(bundle: dict, ext: str = "docx") -> str:
    """报告文件名：患者报告_{编号}_{姓名}_{日期}.{ext}（ext=docx/pdf）"""
    stamp = bundle["generated_at"].strftime("%Y%m%d")
    return f"患者报告_{bundle['patient']['patient_no']}_{bundle['patient']['name']}_{stamp}.{ext}"


# ------------------------- pdf 排版 -------------------------


def _register_cn_font() -> str:
    """注册中文字体供 reportlab 使用；失败时回退默认字体（中文可能缺字）。"""
    candidates = [
        ("MSYH", r"C:\Windows\Fonts\msyh.ttc", 0),
        ("MSYaHei", r"C:\Windows\Fonts\msyh.ttf", None),
        ("SimSun", r"C:\Windows\Fonts\simsun.ttc", 0),
        ("SimHei", r"C:\Windows\Fonts\simhei.ttf", None),
    ]
    for name, path, subfont in candidates:
        try:
            if not os.path.exists(path):
                continue
            if subfont is not None:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=subfont))
            else:
                pdfmetrics.registerFont(TTFont(name, path))
            return name
        except Exception as exc:  # noqa: BLE001
            logger.debug("PDF 字体注册失败 %s: %s", path, exc)
    return "Helvetica"


def _pdf_escape(value) -> str:
    """清洗进入 Paragraph 的文本：多余尖括号可能被当作 XML 标签处理。"""
    return str(value or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_pdf(bundle: dict) -> bytes:
    """把 bundle 渲染为 .pdf 文件字节流（reportlab platypus）。"""
    from reportlab.platypus import KeepTogether, Spacer

    cn = _register_cn_font()

    title_style = ParagraphStyle("title", fontName=cn, fontSize=20, leading=26, alignment=TA_CENTER, spaceAfter=4)
    sub_style = ParagraphStyle("sub", fontName=cn, fontSize=9.5, leading=14, alignment=TA_CENTER, textColor=_MUTED)
    head_style = ParagraphStyle("head", fontName=cn, fontSize=14, leading=18, textColor=_PRIMARY, spaceBefore=12, spaceAfter=6)
    subhead_style = ParagraphStyle("subhead", fontName=cn, fontSize=11.5, leading=15, textColor=_PRIMARY, spaceBefore=10, spaceAfter=4)
    body_style = ParagraphStyle("body", fontName=cn, fontSize=10.5, leading=15, spaceAfter=3)
    muted_style = ParagraphStyle("muted", fontName=cn, fontSize=10.5, leading=15, textColor=_MUTED, spaceAfter=3)
    kv_style = ParagraphStyle("kv", fontName=cn, fontSize=10.5, leading=15, spaceAfter=2)
    cell_style = ParagraphStyle("cell", fontName=cn, fontSize=10, leading=14)
    cell_bold = ParagraphStyle("cellb", fontName=cn, fontSize=10, leading=14)

    def h(text, size=14):
        style = subhead_style if size < 12.5 else head_style
        if size < 12.5:
            return Paragraph(_pdf_escape(text), subhead_style)
        return Paragraph(_pdf_escape(text), head_style)

    story: list = []
    patient = bundle["patient"]

    story.append(Paragraph(_pdf_escape("患者诊疗分析报告"), title_style))
    story.append(
        Paragraph(
            _pdf_escape(
                f"{patient['name']}　{patient['patient_no']}　统计范围：{bundle['range_label']}　"
                f"生成时间：{_fmt_time(bundle['generated_at'])}"
            ),
            sub_style,
        )
    )

    # ---- 一、患者基本信息 ----
    story.append(h("一、患者基本信息"))
    info = patient["info"]
    info_rows = [["", "", "", ""]]
    pair_cols = 4
    for index, (label, value) in enumerate(info):
        row = index // 2
        if len(info_rows) <= row:
            info_rows.append(["", "", "", ""])
        base = (index % 2) * 2
        info_rows[row][base] = Paragraph(_pdf_escape(label), cell_bold)
        info_rows[row][base + 1] = Paragraph(_pdf_escape(value), cell_style)
    if info and len(info) % 2 == 1:
        last = info_rows[-1]
        last[2] = Paragraph("", cell_style)
        last[3] = Paragraph("", cell_style)
    info_table = Table(info_rows, colWidths=[22 * mm, 62 * mm, 22 * mm, 62 * mm], hAlign="LEFT")
    info_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1F5F9")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(info_table)
    if patient.get("primary_diag"):
        story.append(Paragraph(_pdf_escape(f"主诊断：{patient['primary_diag']}"), kv_style))

    # ---- 二、病历记录 ----
    records = bundle["records"]
    story.append(h(f"二、病历记录（{len(records)} 份）"))
    if not records:
        story.append(Paragraph(_pdf_escape("所选时间范围内暂无病历记录。"), muted_style))
    for index, item in enumerate(records, start=1):
        block: list = [Paragraph(_pdf_escape(f"病历 {index} · 记录于 {item['created_at']}"), subhead_style)]
        if not item["fields"]:
            block.append(Paragraph(_pdf_escape("（该病历暂无填写内容）"), muted_style))
        else:
            for label, value in item["fields"]:
                block.append(Paragraph(_pdf_escape(f"{label}：{value}"), kv_style))
        story.append(KeepTogether(block))

    # ---- 三、AI 分析报告 ----
    analyses = bundle["analyses"]
    story.append(h(f"三、AI 分析报告（{len(analyses)} 次）"))
    if not analyses:
        story.append(Paragraph(_pdf_escape("所选时间范围内暂无 AI 分析记录。"), muted_style))
    for index, item in enumerate(analyses, start=1):
        block: list = [Paragraph(_pdf_escape(f"分析 {index} · {item['type_label']} · {item['created_at']}"), subhead_style)]
        block.append(
            Paragraph(
                _pdf_escape(f"模型：{item['model']}　发起人：{item['operator']}　"
                            f"状态：{'成功' if item['status'] == 'done' else '失败'}"),
                muted_style,
            )
        )
        if item["status"] != "done":
            block.append(Paragraph(_pdf_escape(f"失败原因：{item['error'] or '未记录'}"), kv_style))
            story.append(KeepTogether(block))
            continue

        summary = item["summary"] or {}
        if summary:
            block.append(Paragraph(_pdf_escape("综合结论"), body_style))
            for label, value in summary.items():
                block.append(Paragraph(_pdf_escape(f"{label}：{value}"), kv_style))
        else:
            block.append(Paragraph(_pdf_escape("综合结论：（无）"), muted_style))

        attention = item["attention"] or []
        if attention:
            block.append(Paragraph(_pdf_escape("关注点"), body_style))
            for point in attention:
                level = str(point.get("level") or "").strip()
                title = str(point.get("title") or "").strip()
                desc = str(point.get("description") or "").strip()
                label = f"【{level}】{title}" if level else (title or "关注点")
                block.append(Paragraph(_pdf_escape(f"{label}：{desc}"), kv_style))

        evidence = item["evidence"] or []
        if evidence:
            block.append(Paragraph(_pdf_escape("循证依据"), body_style))
            for ev_index, ev in enumerate(evidence, start=1):
                source = str(ev.get("source") or "未命名来源")
                relevance = ev.get("relevance")
                if relevance is not None:
                    source = f"{source}（相关度 {relevance}）"
                block.append(Paragraph(_pdf_escape(f"依据 {ev_index}：{source}"), kv_style))

        story.append(KeepTogether(block))

    story.append(Spacer(1, 18))
    story.append(
        Paragraph(
            _pdf_escape(f"{_DISCLAIMER}　导出人：{bundle['operator']}"),
            ParagraphStyle("footer", fontName=cn, fontSize=9, leading=13, textColor=_MUTED),
        )
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="患者诊疗分析报告",
    )
    doc.build(story)
    return buffer.getvalue()
